"""Create an editable XLSX workbook from the combined manual-review CSV."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_csv")
    parser.add_argument("output_xlsx")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path = Path(args.input_csv)
    output_path = Path(args.output_xlsx)
    with input_path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        first_line = sample.splitlines()[0] if sample else ""
        delimiter = "\t" if "\t" in first_line else ","
        source_rows = list(csv.DictReader(handle, delimiter=delimiter))
    if not source_rows:
        raise ValueError("The manual-review CSV is empty.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Manual Review"
    ws.sheet_view.showGridLines = False

    source_headers = list(source_rows[0])
    manual_headers = [
        "manual_decision",
        "manual_issue",
        "manual_note",
        "corrected_answer",
    ]
    has_manual_columns = all(header in source_headers for header in manual_headers)
    insert_after = source_headers.index("reliability") + 1
    headers = (
        source_headers
        if has_manual_columns
        else source_headers[:insert_after] + manual_headers + source_headers[insert_after:]
    )
    ws.append(headers)
    for row in source_rows:
        values = [row.get(header, "") for header in headers]
        ws.append(values)

    _style_review_sheet(ws, headers)
    _add_summary_sheet(wb, source_rows)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    # Structural QA after serialization.
    check = load_workbook(output_path, read_only=False, data_only=False)
    review = check["Manual Review"]
    assert review.max_row == len(source_rows) + 1
    assert review.max_column == len(headers)
    assert review.auto_filter.ref
    assert "ManualDecisionValidation" in {item.sqref.__str__() and item.errorTitle for item in review.data_validations.dataValidation}
    print(f"Workbook: {output_path}")
    print(f"Rows: {len(source_rows)} | Columns: {len(headers)}")


def _style_review_sheet(ws, headers: list[str]) -> None:
    navy = "203864"
    pale_blue = "D9EAF7"
    editable = "FFF2CC"
    white = "FFFFFF"
    light_border = Side(style="thin", color="D9E1F2")

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="7F8FA6"))
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    widths = {
        "question_id": 11,
        "model": 16,
        "think": 9,
        "question": 42,
        "final_answer": 62,
        "reference_answer": 42,
        "reliability": 16,
        "manual_decision": 18,
        "manual_issue": 28,
        "manual_note": 45,
        "corrected_answer": 55,
        "grounding_issues": 44,
        "expected_tools": 28,
        "acceptable_tools": 28,
        "tool_selection_status": 19,
        "tool_selection_issue": 34,
        "tools_called": 30,
        "tool_args": 42,
        "tool_outputs": 62,
        "tool_call_count": 13,
        "language_ok": 12,
        "language_issue": 28,
        "latency_s": 12,
        "error": 30,
    }
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(header, 18)

    manual_columns = {headers.index(name) + 1 for name in (
        "manual_decision", "manual_issue", "manual_note", "corrected_answer"
    )}
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 72
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=light_border)
            if cell.column in manual_columns:
                cell.fill = PatternFill("solid", fgColor=editable)

    decision_col = get_column_letter(headers.index("manual_decision") + 1)
    validation = DataValidation(
        type="list",
        formula1='"correct,ungrounded,unanswered,unverified_no_tool,incerto,needs_review"',
        allow_blank=True,
    )
    validation.errorTitle = "ManualDecisionValidation"
    validation.error = "Choose correct, ungrounded, unanswered, unverified_no_tool, incerto, or needs_review."
    ws.add_data_validation(validation)
    validation.add(f"{decision_col}2:{decision_col}{ws.max_row}")

    reliability_col = get_column_letter(headers.index("reliability") + 1)
    full_range = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(formula=[f'${reliability_col}2="ungrounded"'], fill=PatternFill("solid", fgColor="FCE4D6")),
    )
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(formula=[f'${reliability_col}2="unanswered"'], fill=PatternFill("solid", fgColor="EDEDED")),
    )
    ws.conditional_formatting.add(
        f"{decision_col}2:{decision_col}{ws.max_row}",
        FormulaRule(formula=[f'${decision_col}2="correct"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )

    table = Table(displayName="ManualReviewTable", ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _add_summary_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Manual Review - Benchmark con CSV"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="203864")
    ws.merge_cells("A1:I1")
    ws["A2"] = "Scena 4 VAL | think=true | run 2026-09-17"
    ws["A4"] = "Record da revisionare"
    ws["B4"] = "=COUNTA('Manual Review'!A:A)-1"
    ws["A5"] = "Domande uniche"
    ws["B5"] = len({row["question_id"] for row in rows})
    ws["A6"] = "Decisioni compilate"
    decision_col = get_column_letter(
        list(wb["Manual Review"].values)[0].index("manual_decision") + 1
    )
    ws["B6"] = f'=COUNTIF(\'Manual Review\'!{decision_col}:{decision_col},"<>")'

    ws.append([])
    ws.append(["Modello", "Record", "Auto ungrounded", "Auto unanswered", "Manual correct", "Manual ungrounded", "Manual unanswered", "Manual unverified", "Manual incerto"])
    models = sorted({row["model"] for row in rows})
    model_col = get_column_letter(list(wb["Manual Review"].values)[0].index("model") + 1)
    reliability_col = get_column_letter(list(wb["Manual Review"].values)[0].index("reliability") + 1)
    for row_index, model in enumerate(models, start=9):
        ws.cell(row_index, 1, model)
        ws.cell(row_index, 2, f'=COUNTIF(\'Manual Review\'!{model_col}:{model_col},A{row_index})')
        ws.cell(row_index, 3, f'=COUNTIFS(\'Manual Review\'!{model_col}:{model_col},A{row_index},\'Manual Review\'!{reliability_col}:{reliability_col},"ungrounded")')
        ws.cell(row_index, 4, f'=COUNTIFS(\'Manual Review\'!{model_col}:{model_col},A{row_index},\'Manual Review\'!{reliability_col}:{reliability_col},"unanswered")')
        ws.cell(row_index, 5, f'=COUNTIFS(\'Manual Review\'!{model_col}:{model_col},A{row_index},\'Manual Review\'!{decision_col}:{decision_col},"correct")')
        ws.cell(row_index, 6, f'=COUNTIFS(\'Manual Review\'!{model_col}:{model_col},A{row_index},\'Manual Review\'!{decision_col}:{decision_col},"ungrounded")')
        ws.cell(row_index, 7, f'=COUNTIFS(\'Manual Review\'!{model_col}:{model_col},A{row_index},\'Manual Review\'!{decision_col}:{decision_col},"unanswered")')
        ws.cell(row_index, 8, f'=COUNTIFS(\'Manual Review\'!{model_col}:{model_col},A{row_index},\'Manual Review\'!{decision_col}:{decision_col},"unverified_no_tool")')
        ws.cell(row_index, 9, f'=COUNTIFS(\'Manual Review\'!{model_col}:{model_col},A{row_index},\'Manual Review\'!{decision_col}:{decision_col},"incerto")')

    for cell in ws[8]:
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(color="FFFFFF", bold=True)
    for col, width in enumerate((20, 14, 18, 18, 18, 20, 20, 20, 16), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A8"


if __name__ == "__main__":
    main()
